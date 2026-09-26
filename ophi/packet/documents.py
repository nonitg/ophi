"""Document renderers: the treatment form (the one place patient identity appears), the narrative as
DOCX and ASCII text, and the single-PDF preview for the on-screen review and sign-off."""

from __future__ import annotations

from datetime import date
from typing import TYPE_CHECKING

from docx import Document
from docx.shared import Inches, Pt
from PIL import Image as PILImage
from reportlab.platypus import Image, KeepTogether, PageBreak
from unidecode import unidecode

from ophi.cdm.models import ArtifactType, AssertionPayload, Case
from ophi.engine.models import Assessment
from ophi.packet import pdf
from ophi.packet.files import PacketFile
from ophi.packet.labels import evidence_label, fmt_date
from ophi.packet.narrative import parse_narrative
from ophi.packet.plates import perio_flowables, psr_flowables
from ophi.rules.schema import RulePack

if TYPE_CHECKING:
    from ophi.service import SignOff

DOC_SPEC = {"dpi": None, "bit_depth": None, "format": "PDF", "colour": "n/a", "pass": True, "transformations": []}


def _fee(cents: int | None) -> str:
    return f"${cents / 100:,.2f}" if cents is not None else ""


# --- treatment form -------------------------------------------------------------------------------


def signature_flowables(case: Case, sign_off: SignOff | None) -> list:
    prov = case.treatment.provider
    licence = f" ({prov.licence})" if prov.licence else ""
    if sign_off:
        so_licence = f" ({sign_off.licence})" if sign_off.licence else ""
        line = (f"Attested electronically by {sign_off.signed_by}{so_licence} at "
                f"{sign_off.signed_at.isoformat()} — {sign_off.attestation}")
    else:
        line = "Signed: ______________________________    Date: ______________"
    return [pdf.spacer(14), pdf.P("Treating provider", "h2"), pdf.P(f"{prov.name}{licence}"), pdf.spacer(4), pdf.P(line)]


def claim_form_flowables(case: Case, sign_off: SignOff | None, prepared_on: date) -> list:
    t = case.treatment
    pt = case.patient
    procedure = [["Code", "Description", "Tooth (FDI)", "Surfaces", "Lab codes", "Fee"],
                 [t.code, t.description or "", str(t.tooth.tooth_fdi), ", ".join(t.surfaces), ", ".join(t.lab_codes), _fee(t.fee_cents)]]
    return [
        pdf.P("Treatment form — CDCP preauthorization request", "title"),
        pdf.P(f"Computer-generated treatment form. Clinic: {case.clinic}. Prepared {fmt_date(prepared_on)}.", "small"),
        pdf.spacer(8),
        pdf.P("Patient", "h2"),
        pdf.grid([["Name", "Date of birth", "CDCP client ID", "PMS patient ID"],
                  [pt.display_name, fmt_date(pt.dob), pt.cdcp_client_id or "not on file", pt.patient_id]],
                 col_widths=[pdf.BODY_WIDTH * f for f in (0.34, 0.2, 0.26, 0.2)]),
        pdf.P("Provider", "h2"),
        pdf.grid([["Name", "Licence", "Clinic"], [t.provider.name, t.provider.licence or "", case.clinic]],
                 col_widths=[pdf.BODY_WIDTH * f for f in (0.34, 0.2, 0.46)]),
        pdf.P("Proposed treatment", "h2"),
        pdf.grid(procedure, col_widths=[pdf.BODY_WIDTH * f for f in (0.1, 0.36, 0.13, 0.13, 0.14, 0.14)]),
        pdf.spacer(4),
        pdf.P(f"Planned date: {fmt_date(t.planned_date)}.    Appointment: {fmt_date(t.appointment_date)}."),
        *signature_flowables(case, sign_off),
        pdf.spacer(10),
        pdf.P("Prepared with Ophi; submitted by the treating provider.", "small"),
    ]


def render_claim_form(case: Case, out_path, sign_off: SignOff | None, prepared_on: date) -> dict:
    pdf.build_pdf(out_path, claim_form_flowables(case, sign_off, prepared_on), "Treatment form")
    return dict(DOC_SPEC)


# --- narrative ------------------------------------------------------------------------------------


def narrative_ascii(text: str) -> str:
    """The spec says ASCII text literally; transliterate rather than emit UTF-8."""
    out = unidecode(text)
    assert out.isascii()
    return out


def render_narrative_docx(text: str, out_path) -> dict:
    doc = Document()
    doc.core_properties.title = "Preauthorization documentation narrative"
    doc.core_properties.author = "Ophi"
    for kind, body in parse_narrative(text):
        if kind == "title":
            doc.add_heading(body, level=0)
        elif kind == "subtitle":
            doc.add_paragraph(body).runs[0].italic = True
        elif kind == "heading":
            doc.add_heading(body, level=1)
        elif kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            p.add_run(body).italic = True
        elif kind == "attribution":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.5)
            run = p.add_run(body)
            run.font.size = Pt(9)
        elif kind == "bullet":
            doc.add_paragraph(body, style="List Bullet")
        else:
            doc.add_paragraph(body)
    doc.save(str(out_path))
    return {**DOC_SPEC, "format": "DOCX"}


def narrative_flowables(text: str) -> list:
    styles = {"title": "h1", "subtitle": "small", "heading": "h2", "quote": "quote", "attribution": "attribution", "body": "body"}
    out = []
    for kind, body in parse_narrative(text):
        if kind == "bullet":
            out.append(pdf.P(f"• {body}"))
        else:
            out.append(pdf.P(body, styles[kind]))
    return out


# --- preview --------------------------------------------------------------------------------------


def _evidence_summary(evidence: list, case: Case) -> str:
    docs = [evidence_label(e, case) for e in evidence if e.type != ArtifactType.CLINICIAN_ASSERTION.value]
    n = len(evidence) - len(docs)
    if n:
        docs.append(f"{n} clinician assertion{'s' if n > 1 else ''} (listed with the narrative)")
    return "; ".join(docs) or "\u2014"


def _completeness_flowables(case: Case, assessment: Assessment) -> list:
    c = assessment.completeness
    rows = [["Requirement", "Status", "Evidence on record"]]
    for r in assessment.requirements:
        if not r.applicable:
            continue
        rows.append([r.label, "SKIPPED (test run)" if r.skipped else r.status.value.replace("_", " "), _evidence_summary(r.evidence, case)])
    out = [
        pdf.P("Documentation completeness", "h1"),
        pdf.P(f"{c['satisfied']} of {c['applicable']} applicable CDCP requirements satisfied under {assessment.ruleset.id} "
              f"{assessment.ruleset.version}, assessed as of {fmt_date(assessment.submission_date_assumed)}. "
              f"Engine status: {assessment.verdict.value}."),
        pdf.spacer(6),
        pdf.grid(rows, col_widths=[pdf.BODY_WIDTH * f for f in (0.42, 0.16, 0.42)]),
    ]
    if assessment.actions:
        out += [pdf.P("Open items", "h2")]
        out += [pdf.P(f"{a.rank}. [{'blocking' if a.blocking else 'advisory'}] {a.title}") for a in assessment.actions]
    return out


def _plate_flowables(f: PacketFile) -> list:
    with PILImage.open(f.path) as im:
        w, h = im.size
    width = pdf.BODY_WIDTH * 0.8
    caption = f.description + (f"  SPEC CHECK FAILED: {f.spec_failure}" if f.spec_failure else "")
    return [KeepTogether([Image(str(f.path), width=width, height=width * h / w), pdf.P(caption, "small"), pdf.spacer(10)])]


def _assertions_flowables(case: Case, pack: RulePack) -> list:
    rows = [["Criterion", "Recorded as", "By", "On"]]
    for a in case.artifacts_of(ArtifactType.CLINICIAN_ASSERTION):
        p = a.payload
        assert isinstance(p, AssertionPayload)
        crit = pack.assertion_criteria.get(p.criterion_id)
        rows.append([crit.label if crit else p.criterion_id, p.value.replace("_", " "),
                     f"{p.asserted_by}" + (f" ({p.licence})" if p.licence else ""), fmt_date(p.asserted_at.date())])
    body = pdf.grid(rows, col_widths=[pdf.BODY_WIDTH * f for f in (0.46, 0.14, 0.26, 0.14)]) if len(rows) > 1 \
        else pdf.P("No clinician assertions recorded.")
    return [pdf.P("Clinician assertions", "h1"), body]


def _index_flowables(files: list[PacketFile]) -> list:
    rows = [["#", "File", "Contents"]] + [[f"{f.seq:02d}", f.filename, f.description] for f in files]
    return [pdf.P("Files in this packet", "h1"), pdf.grid(rows, col_widths=[pdf.BODY_WIDTH * f for f in (0.06, 0.44, 0.5)])]


def render_preview(case: Case, assessment: Assessment, pack: RulePack, narrative_text: str, files: list[PacketFile],
                   sign_off: SignOff | None, out_path) -> None:
    t = case.treatment
    flow = [
        pdf.P("Packet preview — " + ("TEST RUN, chart gaps skipped, not for submission" if assessment.test_run else "review and sign-off"), "title"),
        pdf.P(f"{case.clinic} — reference {assessment.assessment_id} — {t.code} on tooth #{t.tooth.tooth_fdi} — "
              f"{t.provider.name}", "small"),
        pdf.spacer(6),
        *_completeness_flowables(case, assessment),
        PageBreak(),
        *claim_form_flowables(case, sign_off, assessment.submission_date_assumed),
    ]
    plates = [f for f in files if f.kind == "radiograph"]
    if plates:
        flow += [PageBreak(), pdf.P("Radiographs", "h1")]
        for f in plates:
            flow += _plate_flowables(f)
    for f in files:
        art = case.artifact(f.artifact_id) if f.artifact_id else None
        if art and f.kind == "perio_chart":
            flow += [PageBreak(), *perio_flowables(art, case)]
        elif art and f.kind == "psr":
            flow += [PageBreak(), *psr_flowables(art)]
    flow += [PageBreak(), pdf.P("Narrative", "h1"), *narrative_flowables(narrative_text)]
    flow += [pdf.spacer(10), *_assertions_flowables(case, pack)]
    flow += [PageBreak(), *_index_flowables(files), *signature_flowables(case, sign_off)]
    pdf.build_pdf(out_path, flow, "Packet preview")
